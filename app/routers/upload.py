import os,uuid,io
from fastapi import APIRouter, UploadFile, File, HTTPException,Depends
from pydantic import BaseModel
from PyPDF2 import PdfReader
from docx import Document
from app.models.file_record import FileRecord
from app.utils.auth import get_current_user
from openpyxl import load_workbook

router = APIRouter(prefix="/upload", tags=["Upload"])

UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)

class UploadResponse(BaseModel):
    filename: str
    size: int

def extract_text(file_content: bytes, filename: str) -> str:
    """从文件内容中提取文字"""
    ext = filename.lower().split(".")[-1]

    if ext == "pdf":
        reader = PdfReader(io.BytesIO(file_content))
        text = ""
        for page in reader.pages:
            text += page.extract_text() or ""
        return text

    elif ext == "docx":
        doc = Document(io.BytesIO(file_content))
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        return text

    elif ext in ("xlsx", "xls"):
        wb = load_workbook(io.BytesIO(file_content), read_only=True, data_only=True)
        text = ""
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    if cell is not None:
                        text += str(cell) + " "
                text += "\n"
        return text
    
    else:
        raise ValueError(f"不支持的文件格式：.{ext}，目前支持 PDF、Word(.docx)、Excel(.xlsx)")


@router.post("", response_model=UploadResponse)
async def upload(file: UploadFile = File(...),user_id: int = Depends(get_current_user)):
    if not file.filename:
        raise HTTPException(status_code=400, detail="缺少文件名")

    ext = os.path.splitext(file.filename)[1]
    safe_name = f"{uuid.uuid4().hex}{ext}"
    save_path = os.path.join(UPLOAD_DIR, safe_name)

    content = await file.read()
    with open(save_path, "wb") as f:
        f.write(content)

    # 提取文件内容
    try:
        text = extract_text(content, file.filename)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception:
        raise HTTPException(status_code=400, detail="文件解析失败，请确认文件未损坏")

    # 存到数据库
    await FileRecord.create(
        user_id=user_id,
        original_name=file.filename,
        saved_name=safe_name,
        file_size=len(content),
        content=text
    )

    return {"filename": file.filename, "size": len(content)}
